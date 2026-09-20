#!/usr/bin/env python3
# /// script
# dependencies = ["openpyxl"]
# ///
"""Import supplier enzyme data from the comparison xlsx into enzyme_providers.json.

Recognition-sequence columns are intentionally ignored (dirty data); cut info
stays in comm_only_enzymes.json. Name cleaning mirrors the cross-check rules:
strip "FastDigest " prefix, -HF/-HFv2/-v2 suffixes, ", GMP Grade"/", ADCF"
trailing notes, trailing "*"; merge parenthesised aliases and newline-separated
synonyms; fix three known sheet typos.
"""
import json, re, sys, collections
import openpyxl

XLSX = sys.argv[1]
DB_PATH = "backend/libregene-core/data/comm_only_enzymes.json"
OUT_PATH = "backend/libregene-core/data/enzyme_providers.json"

SHEET_PROVIDER = {
    "NEB": "neb",
    "愚公生物_百时美": "bestenzyme",
    "Thermo_FastDigest": "thermo",
}

# canonical alias -> DB name (resolved against the DB at runtime too)
TYPO_FIX = {"Alul": "AluI", "HinP1l": "HinP1I", "Sgel": "SgeI"}


def canonicalize(v):
    v = re.sub(r"-HFv?\d*$", "", v, flags=re.I)
    v = re.sub(r"[- ]v\d+$", "", v, flags=re.I)
    v = re.sub(r",\s*(GMP Grade|ADCF).*$", "", v, flags=re.I).strip()
    return TYPO_FIX.get(v, v)


def name_variants(raw):
    """Return (canonical_names, display_variants).

    canonical_names: fully stripped, used to resolve against the DB.
    display_variants: supplier-facing names worth keeping as aliases
    (e.g. "BbsI-HF", "BsaI-HFv2", "MunI"), with FastDigest prefix and
    trailing "*" removed but variant suffixes intact.
    """
    n = str(raw).strip()
    n = re.sub(r"^FastDigest\s+", "", n)
    n = re.sub(r"\*$", "", n).strip()
    canon, display = [], []
    for part in n.split("\n"):
        part = part.strip()
        if not part:
            continue
        aliases = [a.strip() for a in re.findall(r"\(([^)]+)\)", part)]
        base = re.sub(r"\([^)]*\)", "", part).strip()
        for v in [base] + aliases:
            if not v:
                continue
            canon.append(canonicalize(v))
            dv = TYPO_FIX.get(v, v)
            display.append(dv)
            cd = canonicalize(v)
            if cd != dv:
                canon.append(cd)
    def dedupe(seq):
        seen, res = set(), []
        for v in seq:
            if v.lower() not in seen:
                seen.add(v.lower())
                res.append(v)
        return res
    return dedupe(canon), dedupe(display)


db = json.load(open(DB_PATH))
dbmap = {r["name"].lower(): r["name"] for r in db}


def resolve_db_name(names):
    for n in names:
        if n.lower() in dbmap:
            return dbmap[n.lower()]
    for n in names:
        k = re.sub(r"[^a-z0-9]", "", n.lower())
        for kk, canon in dbmap.items():
            if re.sub(r"[^a-z0-9]", "", kk) == k:
                return canon
    return None


def cell(v):
    if v is None:
        return ""
    return str(v).strip()


def clean_value(key, s):
    """Translate/normalize supplier field values to English."""
    s = re.sub(r"百时美\s*", "", s)  # '百时美 CutOne Buffer' -> 'CutOne Buffer'
    s = s.replace("专用缓冲液，不兼容其他", "N/A (dedicated buffer)")
    m = re.fullmatch(r"最长无星号孵育[:：]([\s\S]+)", s)
    if m:
        times = [t.strip() for t in re.split(r"[\n/]", m.group(1)) if t.strip()]
        times = [t if t.endswith("h") else t + "h" for t in times]
        uniq = list(dict.fromkeys(times))
        span = uniq[0] if len(uniq) == 1 else f"{uniq[0]}\u2013{uniq[-1]}"
        return f"Star-activity-free up to {span}"
    if key == "starActivity" and "% in " in s:
        # '75% in r1.1; 100% in r2.1' -> 'in r1.1, r2.1' (percentages were
        # a data-entry error; the column lists buffers with star activity)
        bufs = []
        for seg in s.split(";"):
            seg = seg.strip()
            mm = re.fullmatch(r"\d+(?:\.\d+)?%\s+in\s+(.+)", seg)
            bufs.append(mm.group(1).strip() if mm else seg)
        return "in " + ", ".join(bufs)
    return s


def row_to_provider_info(r, idx):
    buffers = []
    for bi in range(1, 5):
        name = clean_value("buffers", cell(r[idx[f"b{bi}name"]]))
        act = clean_value("activity", cell(r[idx[f"b{bi}act"]]))
        if name:
            buffers.append({"name": name, "activity": act})
    info = {
        "buffers": buffers,
        "workTemp": cell(r[idx["temp"]]),
        "heatInactivation": cell(r[idx["inact"]]),
        "methylation": cell(r[idx["meth"]]),
        "starActivity": clean_value("starActivity", cell(r[idx["star"]])),
        "catalog": cell(r[idx["cat"]]),
    }
    return {k: v for k, v in info.items() if v != "" and v != []}


wb = openpyxl.load_workbook(XLSX)
enzymes = collections.defaultdict(lambda: {"aliases": [], "providers": {}})
provider_only = collections.defaultdict(lambda: {"aliases": [], "providers": {}})

for sheet, prov in SHEET_PROVIDER.items():
    ws = wb[sheet]
    hdr = [c.value for c in ws[1]]
    idx = {
        "name": hdr.index("酶名称"),
        "temp": hdr.index("工作温度(°C)"),
        "inact": hdr.index("热失活温度"),
        "meth": hdr.index("甲基化影响"),
        "star": hdr.index("星号活性"),
        "cat": hdr.index("货号"),
    }
    for bi in range(1, 5):
        idx[f"b{bi}name"] = hdr.index(f"缓冲液{bi}名称")
        idx[f"b{bi}act"] = hdr.index(f"缓冲液{bi}活性(%)")
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[idx["name"]]:
            continue
        names, display = name_variants(r[idx["name"]])
        if not names:
            continue
        info = row_to_provider_info(r, idx)
        canon = resolve_db_name(names)
        entry = enzymes[canon] if canon else provider_only[display[0]]
        for a in display:
            if a.lower() != (canon or "").lower() and a not in entry["aliases"]:
                entry["aliases"].append(a)
        # Keep each variant's supplier data separate (e.g. BamHI vs BamHI-HF
        # have different buffer compatibilities).
        p = entry["providers"].setdefault(prov, {"variants": {}})
        vname = display[0]
        if vname not in p["variants"]:
            p["variants"][vname] = info

out_enzymes = {}
for canon, e in enzymes.items():
    provs = {}
    for pk, p in e["providers"].items():
        variants = [{"name": vn, **vi} for vn, vi in p["variants"].items()]
        # canonical-named variant first
        variants.sort(key=lambda v: 0 if v["name"].lower() == canon.lower() else 1)
        provs[pk] = {"variants": variants}
    out_enzymes[canon] = {"aliases": e["aliases"], "providers": provs}
def variants_of(e):
    provs = {}
    for pk, p in e["providers"].items():
        provs[pk] = {"variants": [{"name": vn, **vi} for vn, vi in p["variants"].items()]}
    return provs


out_only = [
    {"name": name, "aliases": e["aliases"], "providers": variants_of(e)}
    for name, e in sorted(provider_only.items())
]

result = {"enzymes": dict(sorted(out_enzymes.items())), "providerOnly": out_only}
with open(OUT_PATH, "w") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)
    f.write("\n")

print("enzymes:", len(out_enzymes))
print("providerOnly:", len(out_only), [e["name"] for e in out_only])
per_prov = collections.Counter(
    p for e in out_enzymes.values() for p in e["providers"]
)
per_prov.update(p for e in out_only for p in e["providers"])
print("per provider:", dict(per_prov))
